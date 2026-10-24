"""Readable copies of the patient backup: an Excel workbook and a Word document.

These are for reading and printing only. They are built from the same data as
the .json backup (see patient_backup.py) but cannot be used to restore.

Needs:  pip install openpyxl python-docx
"""

import io
from datetime import date, datetime, timedelta, timezone

PH_TZ = timezone(timedelta(hours=8))
CLINIC = "Capizonda Dental Clinic"
BLUE = "0B4AA2"
LIGHT = "EAF1FB"


# ---------------------------------------------------------------------------
# 1. Turn a backup payload into plain, display-ready records
# ---------------------------------------------------------------------------
def _num(v):
    try:
        n = float(v)
        return n if n == n and n not in (float("inf"), float("-inf")) else 0.0
    except Exception:
        return 0.0


def _text(v):
    if v is None:
        return ""
    if isinstance(v, dict) and v.get("__fs__") == "datetime":
        return str(v.get("v", ""))[:16].replace("T", " ")
    return str(v).strip()


def _age(birthday, today):
    try:
        b = date.fromisoformat(str(birthday)[:10])
    except Exception:
        return ""
    return today.year - b.year - ((today.month, today.day) < (b.month, b.day))


def _display_name(first, middle, last):
    given = " ".join(x for x in (first, middle) if x)
    if last and given:
        return f"{last}, {given}"
    return last or given


def build_records(payload):
    """-> (patients, meta). Each patient is a dict with name, details,
    treatments, payments and appointments, all as plain text/numbers."""
    today = datetime.now(PH_TZ).date()
    accounts = {a["id"]: a for a in payload.get("accounts", [])}
    # Same rule as the My Patients table: one row per ACCOUNT, and the first
    # Patients document (lowest Patient ID) pointing at it supplies the ID.
    by_uid = {}
    for p in sorted(payload.get("patients", []), key=lambda x: x["id"]):
        uid = _text((p.get("data") or {}).get("account_uid"))
        if uid and uid not in by_uid:
            by_uid[uid] = p

    def make(patient, account):
        pd = (patient or {}).get("data") or {}
        ad = (account or {}).get("data") or {}
        # the account's name parts are the editable ones the admin sees
        first = _text(ad.get("firstname") or pd.get("first_name"))
        middle = _text(ad.get("middlename") or pd.get("middle_name"))
        last = _text(ad.get("lastname") or pd.get("last_name"))
        name = _display_name(first, middle, last) or _text(ad.get("name")) or "(no name)"
        birthday = _text(pd.get("birthday") or ad.get("birthday"))
        provider = _text(ad.get("provider"))
        rec = {
            "patient_id": _text((patient or {}).get("id")),
            "name": name,
            "birthday": birthday,
            "age": _age(birthday, today),
            "sex": _text(ad.get("last_sex") or ad.get("sex") or pd.get("sex")),
            "civil": _text(ad.get("CivilStatus") or ad.get("civil_status") or pd.get("civil_status")),
            "mobile": _text(ad.get("contact_number") or pd.get("contact_number")),
            "email": _text(ad.get("email") or pd.get("email")),
            "account": {"walk_in": "Walk-in", "google": "Google", "password": "Email"}.get(provider, provider),
            "treatments": [], "payments": [], "appointments": [],
        }
        subs = (account or {}).get("subcollections", {})
        for doc in subs.get("Done_procedure", []):
            for row in (doc.get("data") or {}).get("procedures") or []:
                if not isinstance(row, dict):
                    continue
                t = {
                    "date": _text(row.get("date")),
                    "procedure": _text(row.get("procedure")),
                    "tooth": _text(row.get("tooth")),
                    "dentist": _text(row.get("dentist")),
                    "medicine": _text(row.get("medicine")),
                    "fee": _num(row.get("value")),
                    "paid": _num(row.get("paid")),
                    "balance": _num(row.get("balance")),
                    "status": _text(row.get("status")),
                    "next": _text(row.get("next_appointment")),
                }
                rec["treatments"].append(t)
                for pay in row.get("payments") or []:
                    if isinstance(pay, dict):
                        rec["payments"].append({
                            "date": _text(pay.get("date")),
                            "procedure": t["procedure"],
                            "method": _text(pay.get("method")) or "Not recorded",
                            "amount": _num(pay.get("amount")),
                        })
        for sub in ("Approve", "appointments"):
            for doc in subs.get(sub, []):
                d = doc.get("data") or {}
                when = _text(d.get("appointment_date")) or " ".join(
                    x for x in (_text(d.get("date")), _text(d.get("time"))) if x)
                rec["appointments"].append({
                    "when": when,
                    "service": _text(d.get("Service") or d.get("service")),
                    "dentist": _text(d.get("DentistName") or d.get("dentist_name") or d.get("dentist")),
                    "status": _text(d.get("status")),
                })
        rec["treatments"].sort(key=lambda r: r["date"])
        rec["payments"].sort(key=lambda r: r["date"])
        rec["appointments"].sort(key=lambda r: r["when"])
        rec["total_fee"] = sum(r["fee"] for r in rec["treatments"])
        rec["total_paid"] = sum(r["paid"] for r in rec["treatments"])
        rec["total_balance"] = sum(r["balance"] for r in rec["treatments"])
        return rec

    patients = []
    for uid, acc in accounts.items():
        patient = by_uid.get(uid)
        patients.append(make(patient, acc))   # every account is listed
    for p in sorted(payload.get("patients", []), key=lambda x: x["id"]):
        uid = _text((p.get("data") or {}).get("account_uid"))
        if not uid or uid not in accounts:      # a patient with no account at all
            patients.append(make(p, None))
    patients.sort(key=lambda r: r["name"].lower())

    meta = {
        "generated": datetime.now(PH_TZ).strftime("%B %d, %Y %I:%M %p").replace(" 0", " "),
        "patients": len(patients),
        "with_records": sum(1 for p in patients if p["treatments"]),
        "treatments": sum(len(p["treatments"]) for p in patients),
        "payments": sum(len(p["payments"]) for p in patients),
        "fee": sum(p["total_fee"] for p in patients),
        "paid": sum(p["total_paid"] for p in patients),
        "balance": sum(p["total_balance"] for p in patients),
    }
    return patients, meta


def peso(n):
    return "\u20b1{:,.2f}".format(n or 0)


# ---------------------------------------------------------------------------
# 2. Excel
# ---------------------------------------------------------------------------
def build_xlsx(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    patients, meta = build_records(payload)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=BLUE)
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D5DEEE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '"\u20b1"#,##0.00'

    def sheet(ws, headers, rows, widths, money_cols=()):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if c.column in money_cols:
                    c.number_format = money
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = CLINIC
    ws["A1"].font = Font(bold=True, size=16, color=BLUE)
    ws["A2"] = "Patient records report"
    ws["A3"] = "Generated: " + meta["generated"]
    ws["A4"] = "Confidential patient information. Keep this file somewhere safe."
    ws["A4"].font = Font(italic=True, color="B91C1C")
    lines = [
        ("Patients", meta["patients"]), ("Patients with treatment records", meta["with_records"]),
        ("Treatment records", meta["treatments"]), ("Payments recorded", meta["payments"]),
        ("Total fees", meta["fee"]), ("Total paid", meta["paid"]), ("Total balance", meta["balance"]),
    ]
    for i, (k, v) in enumerate(lines, start=6):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        c.alignment = Alignment(horizontal="right")
        if k.startswith("Total"):
            c.number_format = money
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 18

    # Patients
    sheet(wb.create_sheet("Patients"),
          ["Patient ID", "Name", "Age", "Birthday", "Sex", "Civil Status", "Mobile", "Email", "Account",
           "Treatments", "Total Fees", "Total Paid", "Balance"],
          [[p["patient_id"], p["name"], p["age"], p["birthday"], p["sex"], p["civil"], p["mobile"],
            p["email"], p["account"], len(p["treatments"]), p["total_fee"], p["total_paid"],
            p["total_balance"]] for p in patients],
          [12, 30, 7, 12, 9, 12, 16, 30, 10, 11, 14, 14, 14], money_cols=(11, 12, 13))

    # Treatments
    rows = []
    for p in patients:
        for t in p["treatments"]:
            rows.append([p["name"], p["patient_id"], t["date"], t["procedure"], t["tooth"], t["dentist"],
                         t["medicine"], t["fee"], t["paid"], t["balance"], t["status"], t["next"]])
    sheet(wb.create_sheet("Treatments"),
          ["Patient", "Patient ID", "Date", "Procedure", "Tooth", "Dentist", "Medicine", "Fee", "Paid",
           "Balance", "Status", "Next Appointment"],
          rows, [30, 12, 12, 26, 9, 26, 18, 14, 14, 14, 10, 18], money_cols=(8, 9, 10))

    # Payments
    rows = [[p["name"], p["patient_id"], y["date"], y["procedure"], y["method"], y["amount"]]
            for p in patients for y in p["payments"]]
    sheet(wb.create_sheet("Payments"),
          ["Patient", "Patient ID", "Date", "Procedure", "Method", "Amount"],
          rows, [30, 12, 12, 28, 16, 14], money_cols=(6,))

    # Appointments
    rows = [[p["name"], p["patient_id"], a["when"], a["service"], a["dentist"], a["status"]]
            for p in patients for a in p["appointments"]]
    sheet(wb.create_sheet("Appointments"),
          ["Patient", "Patient ID", "Date & Time", "Service", "Dentist", "Status"],
          rows, [30, 12, 18, 28, 26, 14])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# 3. Word
# ---------------------------------------------------------------------------
def build_docx(payload):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    patients, meta = build_records(payload)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name, st.font.size = "Calibri", Pt(size)
        st.font.bold, st.font.color.rgb = True, RGBColor.from_string(BLUE)

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def table(headers, rows, widths, right_cols=(), bold_last=False):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(255, 255, 255)
            shade(c, BLUE)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(9)
                if bold_last and ri == len(rows) - 1:
                    run.bold = True
                    shade(cells[i], LIGHT)
                if i in right_cols:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)          # grid widths (LibreOffice, Word)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        # repeat header row on each page
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)
        return t

    # footer: confidential + page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(f"{CLINIC}  \u00b7  Confidential patient records  \u00b7  Page ")
    run.font.size = Pt(8)
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            r2._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r2._r.append(it)

    # cover / summary
    h = doc.add_heading(CLINIC, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph("Patient Records Report").runs[0].font.size = Pt(13)
    doc.add_paragraph("Generated: " + meta["generated"])
    warn = doc.add_paragraph()
    wr = warn.add_run("Confidential patient information. Keep this document somewhere safe and do not share it.")
    wr.italic, wr.font.color.rgb = True, RGBColor(0xB9, 0x1C, 0x1C)
    table(["Summary", ""], [
        ["Patients", meta["patients"]], ["Patients with treatment records", meta["with_records"]],
        ["Treatment records", meta["treatments"]], ["Payments recorded", meta["payments"]],
        ["Total fees", peso(meta["fee"])], ["Total paid", peso(meta["paid"])],
        ["Total balance", peso(meta["balance"])],
    ], [9, 5], right_cols=(1,))

    # index of patients
    doc.add_heading("Patients", level=2)
    table(["Patient ID", "Name", "Treatments", "Total Fees", "Total Paid", "Balance"],
          [[p["patient_id"] or "-", p["name"], len(p["treatments"]), peso(p["total_fee"]),
            peso(p["total_paid"]), peso(p["total_balance"])] for p in patients],
          [3, 9, 3, 4, 4, 4], right_cols=(2, 3, 4, 5))

    # one section per patient
    for p in patients:
        doc.add_page_break()
        doc.add_heading(p["name"], level=2)
        info = [("Patient ID", p["patient_id"] or "-"),
                ("Birthday / Age", (p["birthday"] or "-") + (f"  ({p['age']} yrs)" if p["age"] != "" else "")),
                ("Sex", p["sex"] or "-"), ("Civil Status", p["civil"] or "-"),
                ("Mobile", p["mobile"] or "-"), ("Email", p["email"] or "-"),
                ("Account", p["account"] or "-")]
        t = doc.add_table(rows=0, cols=4)
        t.style = "Table Grid"
        t.autofit = False
        for i in range(0, len(info), 2):
            cells = t.add_row().cells
            for j, (k, v) in enumerate(info[i:i + 2]):
                cells[j * 2].text = ""
                cells[j * 2].paragraphs[0].add_run(k).bold = True
                shade(cells[j * 2], LIGHT)
                cells[j * 2 + 1].text = str(v)
        for i, w in enumerate((3.2, 9.5, 3.2, 9.5)):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate((3.2, 9.5, 3.2, 9.5)):
                row.cells[i].width = Cm(w)

        doc.add_heading("Treatment history", level=3)
        if p["treatments"]:
            rows = [[r["date"], r["procedure"], r["tooth"], r["dentist"], peso(r["fee"]), peso(r["paid"]),
                     peso(r["balance"]), r["status"]] for r in p["treatments"]]
            rows.append(["Totals", "", "", "", peso(p["total_fee"]), peso(p["total_paid"]),
                         peso(p["total_balance"]), ""])
            table(["Date", "Procedure", "Tooth", "Dentist", "Fee", "Paid", "Balance", "Status"],
                  rows, [2.6, 5.4, 2.2, 5.2, 3, 3, 3, 2.2], right_cols=(4, 5, 6), bold_last=True)
        else:
            doc.add_paragraph("No treatment records.")

        if p["payments"]:
            doc.add_heading("Payments", level=3)
            table(["Date", "Procedure", "Method", "Amount"],
                  [[y["date"], y["procedure"], y["method"], peso(y["amount"])] for y in p["payments"]],
                  [3, 8, 4, 4], right_cols=(3,))

        if p["appointments"]:
            doc.add_heading("Appointments", level=3)
            table(["Date & Time", "Service", "Dentist", "Status"],
                  [[a["when"], a["service"], a["dentist"], a["status"]] for a in p["appointments"]],
                  [4, 8, 7, 3])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()