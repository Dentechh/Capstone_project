import io, json, sys, copy
from datetime import datetime, UTC
sys.path.insert(0, ".")
import patient_backup as pb
from flask import Flask, session

# ---------- tiny in-memory Firestore ----------
class Snap:
    def __init__(self, ref, data): self.reference, self.id, self._d = ref, ref.id, data
    exists = property(lambda s: s._d is not None)
    def to_dict(self): return copy.deepcopy(self._d) if self._d is not None else None
class Doc(pb._DocumentReference):
    def __init__(self, db, path): self.db, self.path, self.id = db, path, path.split("/")[-1]
    @property
    def reference(self): return self
    def get(self, transaction=None): return Snap(self, self.db.store.get(self.path))
    def set(self, data, merge=False):
        cur = self.db.store.get(self.path, {}) if merge else {}
        self.db.store[self.path] = {**cur, **copy.deepcopy(data)}
    def collection(self, name): return Col(self.db, f"{self.path}/{name}")
class Col:
    def __init__(self, db, path): self.db, self.path = db, path
    def document(self, i): return Doc(self.db, f"{self.path}/{i}")
    def stream(self):
        n = self.path.count("/") + 2
        for p in sorted(self.db.store):
            if p.startswith(self.path + "/") and p.count("/") + 1 == n:
                yield Snap(Doc(self.db, p), self.db.store[p])
class Batch:
    def __init__(self, db): self.db, self.ops = db, []
    def set(self, ref, data): self.ops.append((ref, data))
    def commit(self):
        self.db.commits += 1
        if self.db.fail_on_commit == self.db.commits: raise RuntimeError("boom")
        for r, d in self.ops: r.set(d)
class DB:
    def __init__(self): self.store, self.commits, self.fail_on_commit = {}, 0, None
    def collection(self, n): return Col(self, n)
    def document(self, p): return Doc(self, p)
    def batch(self): return Batch(self)
    def get_all(self, refs): return [r.get() for r in refs]

class FakeApp(pb.BackupMixin):
    Doc_Patients, Customer_Account, Appointment_cliets = "Patients", "Customer_Account", "appointments"
    def __init__(self, db):
        self.db = db; self.app = Flask(__name__); self.app.secret_key = "x"
        self.app.config["TESTING"] = True
        self.agg_rows = []; self.inval = 0
        self._register_backup_routes()
        @self.app.route("/login-admin")
        def la(): session["admin_logged_in"] = True; session["admin_name"] = "Dr. Test"; return "ok"
    def _invalidate_accounts_cache(self): self.inval += 1
    def _invalidate_patients_cache(self): self.inval += 1
    def _invalidate_financial_cache(self): self.inval += 1
    def _agg_apply_rows(self, add=None, remove=None): self.agg_rows += add or []

def seed(db):
    t = datetime(2026, 10, 1, 8, 30, tzinfo=UTC)
    db.store["Patients/P-000007"] = {"first_name": "Felix", "account_uid": "uid1", "birthday": "2005-01-06", "created_at": "x"}
    db.store["Patients/P-000012"] = {"first_name": "Walk", "account_uid": "walkin_abc"}
    db.store["Customer_Account/uid1"] = {"firstname": "Felix", "has_history": True, "last_approved_at": t, "provider": "password"}
    db.store["Customer_Account/walkin_abc"] = {"firstname": "Walk", "provider": "walk_in"}
    db.store["Customer_Account/uid1/Done_procedure/d1"] = {"updated_at": t, "chart_image": "data:image/png;base64,AAAA",
        "procedures": [{"procedure": "Cleaning", "paid": 500, "balance": 0, "value": 500, "payments": [{"amount": 500, "date": "2026-10-01"}]},
                       {"procedure": "Pasta", "paid": 100, "balance": 400, "value": 500}]}
    db.store["Customer_Account/uid1/Approve/a1"] = {"accepted_at": "2026-10-01", "raw": b"\x00\x01"}
    db.store["Customer_Account/uid1/appointments/p1"] = {"date": "2026-10-02"}
    db.store["Counters/patients"] = {"last_number": 12}

def client(app, admin=True):
    c = app.app.test_client()
    if admin: c.get("/login-admin")
    return c
def upload(c, path, raw, **extra):
    return c.post(path, data={"backup_file": (io.BytesIO(raw), "b.json"), **extra}, content_type="multipart/form-data")

ok = 0
def check(cond, msg):
    global ok
    assert cond, "FAIL: " + msg
    ok += 1; print("  ok -", msg)

# 1 auth
db = DB(); seed(db); app = FakeApp(db)
for p in ("export", "preview", "import"):
    r = client(app, admin=False).post(f"/admin/backup/{p}")
    check(r.status_code == 403, f"{p} rejects non-admin")

# 2 export
r = client(app).post("/admin/backup/export")
check(r.status_code == 200 and "attachment" in r.headers["Content-Disposition"], "export downloads a file")
raw = r.data; bk = json.loads(raw)
check(bk["counts"] == {"patients": 2, "accounts": 2, "treatment_records": 2}, "counts right")
check(app._backup_validate(bk) is None, "export validates")
acc = {a["id"]: a for a in bk["accounts"]}
check(set(acc["uid1"]["subcollections"]) == {"Done_procedure", "Approve", "appointments"}, "all 3 record types exported")
check("chart_image" in acc["uid1"]["subcollections"]["Done_procedure"][0]["data"], "chart included by default")
nochart = json.loads(client(app).post("/admin/backup/export", data={"include_charts": "0"}).data)
check("chart_image" not in {a["id"]: a for a in nochart["accounts"]}["uid1"]["subcollections"]["Done_procedure"][0]["data"], "charts can be left out")
check(app._backup_validate(nochart) is None, "no-chart export validates")

# 3 restore into empty DB == identical data
db2 = DB(); app2 = FakeApp(db2); c2 = client(app2)
pv = upload(c2, "/admin/backup/preview", raw).get_json()["summary"]
check(pv["patients_new"] == 2 and pv["accounts_new"] == 2 and pv["treatment_records_new"] == 2, "preview counts new items")
check(db2.store == {}, "preview wrote nothing")
res = upload(c2, "/admin/backup/import", raw).get_json()
check(res["success"], "import succeeds")
exp = copy.deepcopy(db.store)
check(db2.store == exp, "restored database is identical (dates, bytes, nested payments)")
check(len(app2.agg_rows) == 2 and app2.inval == 3, "aggregates updated + caches cleared")

# 4 second import skips everything, overwrites nothing
db2.store["Patients/P-000007"]["first_name"] = "CHANGED"
res = upload(c2, "/admin/backup/import", raw).get_json()
check(res["summary"]["patients_new"] == 0 and res["summary"]["accounts_new"] == 0, "re-import skips existing")
check(db2.store["Patients/P-000007"]["first_name"] == "CHANGED", "existing data never overwritten")
check(len(app2.agg_rows) == 2, "no double counting in aggregates on re-import")

# 5 partial: only the deleted patient is restored
for k in [k for k in db2.store if "uid1" in k]: del db2.store[k]
del db2.store["Patients/P-000007"]
res = upload(c2, "/admin/backup/import", raw).get_json()
check(res["summary"]["accounts_new"] == 1 and res["summary"]["patients_new"] == 1, "only missing patient/account restored")
check(db2.store == exp, "deleted patient fully restored with transactions")

# 6 counter bump
db3 = DB(); app3 = FakeApp(db3); db3.store["Counters/patients"] = {"last_number": 3}
upload(client(app3), "/admin/backup/import", raw)
check(db3.store["Counters/patients"]["last_number"] == 12, "patient ID counter moved up")

# 7 tampering / garbage
t = json.loads(raw); t["accounts"][0]["data"]["firstname"] = "Evil"
r = upload(client(FakeApp(DB())), "/admin/backup/import", json.dumps(t).encode())
check(r.status_code == 400 and "changed or damaged" in r.get_json()["message"], "edited file rejected")
for bad, label in ((b"not json", "garbage"), (b"{}", "wrong format"),
                   (json.dumps({**bk, "accounts": [{"id": "../x", "data": {}}]}).encode(), "bad id")):
    r = upload(client(FakeApp(DB())), "/admin/backup/import", bad)
    check(r.status_code == 400, f"{label} rejected")
t = json.loads(raw); t["accounts"][0]["subcollections"]["Admin"] = []
r = upload(client(FakeApp(DB())), "/admin/backup/import", json.dumps(t).encode())
check(r.status_code == 400, "unknown record type rejected")
r = client(FakeApp(DB())).post("/admin/backup/import", data={})
check(r.status_code == 400, "missing file rejected")

# 8 interrupted restore can be finished by running again
db4 = DB(); app4 = FakeApp(db4); db4.fail_on_commit = 1
orig = pb._BatchWriter.__init__
pb._BatchWriter.__init__ = lambda self, db, **k: orig(self, db, max_ops=2)
r = upload(client(app4), "/admin/backup/import", raw)
check(r.status_code == 500, "failure reported")
db4.fail_on_commit = None
r = upload(client(app4), "/admin/backup/import", raw).get_json()
pb._BatchWriter.__init__ = orig
check(r["success"] and db4.store == exp, "second run finishes an interrupted restore completely")

# 9 readable copies (Excel / Word)
import openpyxl, docx, importlib
db5 = DB(); seed(db5); app5 = FakeApp(db5); c5 = client(app5)
check(client(app5, admin=False).post("/admin/backup/report", data={"format": "xlsx"}).status_code == 403, "report rejects non-admin")
check(c5.post("/admin/backup/report", data={"format": "pdf"}).status_code == 400, "unknown format rejected")
r = c5.post("/admin/backup/report", data={"format": "xlsx"})
check(r.status_code == 200 and ".xlsx" in r.headers["Content-Disposition"], "Excel downloads")
wb = openpyxl.load_workbook(io.BytesIO(r.data))
check(wb.sheetnames == ["Summary", "Patients", "Treatments", "Payments", "Appointments"], "Excel has the 5 sheets")
check(wb["Treatments"].max_row - 1 == 2 and wb["Payments"].max_row - 1 == 1, "Excel treatment/payment rows right")
check(wb["Patients"]["B2"].value and wb["Patients"].max_row - 1 == 2, "Excel lists both accounts")
r = c5.post("/admin/backup/report", data={"format": "docx"})
check(r.status_code == 200 and ".docx" in r.headers["Content-Disposition"], "Word downloads")
d = docx.Document(io.BytesIO(r.data))
check(any("Felix" in p.text for p in d.paragraphs), "Word shows the patient")
check("chart_image" not in r.data.decode("latin-1"), "Word has no chart images")
sys.modules["patient_report"] = None
check(c5.post("/admin/backup/report", data={"format": "docx"}).status_code == 501, "missing packages give a clear message")
del sys.modules["patient_report"]
print(f"\nALL {ok} CHECKS PASSED")